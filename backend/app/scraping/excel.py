from datetime import datetime

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def save_to_excel(rows, filepath):
    """Save results to a nicely formatted Excel spreadsheet."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Docket Prospects'

    BLUE       = '2563EB'
    WHITE      = 'FFFFFF'
    LIGHT_BLUE = 'EFF6FF'
    GREY_ROW   = 'F8FAFC'

    ws.merge_cells('A1:H1')
    title_cell = ws['A1']
    title_cell.value = f'Docket Prospect List — Generated {datetime.now().strftime("%d %b %Y %H:%M")}'
    title_cell.font      = Font(bold=True, color=WHITE, size=13, name='Calibri')
    title_cell.fill      = PatternFill('solid', fgColor=BLUE)
    title_cell.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 28

    headers = ['Business Name', 'Email', 'Phone Number',
               'Website', 'Address', 'City', 'Category', 'Source']
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=2, column=col, value=h)
        c.font      = Font(bold=True, color=WHITE, size=11, name='Calibri')
        c.fill      = PatternFill('solid', fgColor='1D4ED8')
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws.row_dimensions[2].height = 22

    thin = Side(style='thin', color='E2E8F0')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for i, row in enumerate(rows, 3):
        fill_color = LIGHT_BLUE if i % 2 == 0 else GREY_ROW
        fill = PatternFill('solid', fgColor=fill_color)

        values = [
            row.get('Business Name', ''),
            row.get('Email', ''),
            row.get('Phone Number', ''),
            row.get('Website', ''),
            row.get('Address', ''),
            row.get('City', ''),
            row.get('Category', ''),
            row.get('Source', ''),
        ]
        for col, val in enumerate(values, 1):
            c = ws.cell(row=i, column=col, value=val)
            c.font      = Font(size=10, name='Calibri')
            c.fill      = fill
            c.border    = border
            c.alignment = Alignment(vertical='center', wrap_text=False)

        ws.row_dimensions[i].height = 18

    widths = [35, 35, 20, 40, 45, 15, 22, 16]
    for col, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = width

    ws.freeze_panes = 'A3'

    ws2 = wb.create_sheet(title='Summary')
    ws2['A1'] = 'Summary'
    ws2['A1'].font = Font(bold=True, size=14, color=BLUE, name='Calibri')

    total      = len(rows)
    with_email = sum(1 for r in rows if r.get('Email'))
    with_phone = sum(1 for r in rows if r.get('Phone Number'))

    summary_rows = [
        ('Total businesses found', total),
        ('With email address',     with_email),
        ('With phone number',      with_phone),
        ('Email coverage',         f'{round(with_email / total * 100)}%' if total else '0%'),
        ('Phone coverage',         f'{round(with_phone / total * 100)}%' if total else '0%'),
    ]
    for i, (label, value) in enumerate(summary_rows, 3):
        ws2.cell(row=i, column=1, value=label).font = Font(bold=True, name='Calibri')
        ws2.cell(row=i, column=2, value=value).font = Font(name='Calibri')

    ws2.column_dimensions['A'].width = 30
    ws2.column_dimensions['B'].width = 20

    wb.save(filepath)
